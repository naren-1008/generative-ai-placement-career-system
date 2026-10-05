import os
import logging
import fitz  # PyMuPDF
import docx

logger = logging.getLogger(__name__)

class ResumeParserService:
    @staticmethod
    def extract_text(file_path):
        """
        Extracts raw plain text from a .pdf or .docx file path.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found at: {file_path}")

        ext = file_path.rsplit('.', 1)[-1].lower()

        if ext == 'pdf':
            return ResumeParserService._extract_pdf_text(file_path)
        elif ext == 'docx':
            return ResumeParserService._extract_docx_text(file_path)
        else:
            raise ValueError(f"Unsupported file format: .{ext}. Only PDF and DOCX files are allowed.")

    @staticmethod
    def _extract_pdf_text(file_path):
        text_content = []
        try:
            doc = fitz.open(file_path)
            for page in doc:
                text_content.append(page.get_text())
            doc.close()
            return "\n".join(text_content).strip()
        except Exception as e:
            logger.error(f"Error reading PDF file {file_path}: {e}")
            raise RuntimeError(f"Failed to parse PDF document: {str(e)}")

    @staticmethod
    def _extract_docx_text(file_path):
        try:
            doc = docx.Document(file_path)
            full_text = [para.text for para in doc.paragraphs if para.text.strip()]
            
            # Also extract text from tables if present
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        full_text.append(row_text)
                        
            return "\n".join(full_text).strip()
        except Exception as e:
            logger.error(f"Error reading DOCX file {file_path}: {e}")
            raise RuntimeError(f"Failed to parse DOCX document: {str(e)}")
